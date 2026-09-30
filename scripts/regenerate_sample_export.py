"""
Regenerate the sample XLSX/CSV/JSON export with valid Phase 5 scoring.

Usage:
    python scripts/regenerate_sample_export.py
"""
import os
import sys
import shutil

# Ensure project root is in path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.schemas import (
    CompanyLead, EnrichedLead, ContactCandidate, CompanyEnrichment, StructuredJob
)
from sales_engine.exports.schemas import ExportRequest
from sales_engine.exports.export_orchestrator import ExportOrchestrator

SAMPLE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "exports", "sample_export")


def build_sample_leads():
    """Build sample leads using valid Phase 5 scoring contract."""

    # Lead 1: Qualified high-intent BIM automation lead
    # Phase 5 contractual bounds: fit<=30, intent<=30, recency<=20, evidence<=20, threshold=60
    l1 = CompanyLead(
        company_name="VDC Automation Lab GmbH",
        company_name_normalized="vdc automation lab",
        company_domain="vdc-autolab.de",
        source_company_keys=["vdc-autolab"],
        source_company_identities=["lever:vdc-autolab"],
        job_count=4,
        total_job_count=4,
        unique_job_count=4,
        relevant_job_count=3,
        job_titles=["Head of Digital Practice", "Revit API Developer", "Dynamo Specialist"],
        locations=["Munich, Germany", "Remote"],
        technologies=["Revit", "Dynamo", "C#", "Python"],
        signals=[{"signal": "revit_api", "evidence_count": 3}],
        evidence=["Develop Revit C# addins", "Automate QA with Dynamo"],
        fit_score=30,       # <= 30 ✓
        intent_score=30,    # <= 30 ✓
        recency_score=20,   # <= 20 ✓
        evidence_score=14,  # <= 20 ✓
        lead_score=94,      # 30 + 30 + 20 + 14 = 94 ✓
        qualification_threshold=60,
        qualified=True,
        scoring_reasons=["High fit BIM role", "Active Revit automation hiring"],
    )
    c1_1 = ContactCandidate(
        first_name="Maximilian",
        last_name="Bauer",
        full_name="Maximilian Bauer",
        job_title="Head of BIM & Digital Delivery",
        seniority="Head",
        department="Digital Practice",
        company_name="VDC Automation Lab GmbH",
        company_domain="vdc-autolab.de",
        work_email="m.bauer@vdc-autolab.de",
        email_status="verified",
        email_confidence=98,
        linkedin_url="https://linkedin.com/in/max-bauer-bim",
        buyer_role_match="exact",
        contact_score=95,
        data_sources=["apollo"],
        is_best_contact=True,
    )
    c1_2 = ContactCandidate(
        first_name="Anna",
        last_name="Schmidt",
        full_name="Anna Schmidt",
        job_title="Lead BIM Software Engineer",
        seniority="Lead",
        department="Software",
        company_name="VDC Automation Lab GmbH",
        company_domain="vdc-autolab.de",
        work_email="a.schmidt@vdc-autolab.de",
        email_status="likely",
        email_confidence=80,
        linkedin_url="https://linkedin.com/in/anna-schmidt-eng",
        buyer_role_match="strong",
        contact_score=85,
        data_sources=["hunter"],
        is_best_contact=False,
    )
    e1 = EnrichedLead(
        base_lead=l1,
        company_enrichment=CompanyEnrichment(
            company_name="VDC Automation Lab GmbH",
            company_domain="vdc-autolab.de",
            industry="Architecture & Planning",
            employee_count=120,
            country="Germany",
        ),
        buyer_roles_searched=["Head of BIM", "BIM Director"],
        contacts=[c1_1, c1_2],
        best_contact=c1_1,
        enrichment_status="complete",
        providers_used=["apollo", "hunter"],
        enrichment_errors=[],
    )

    # Lead 2: Moderate qualified lead
    l2 = CompanyLead(
        company_name="Nordic BIM Consulting AS",
        company_name_normalized="nordic bim consulting",
        company_domain="nordicbim.no",
        source_company_keys=["nordic-bim"],
        source_company_identities=["greenhouse:nordic-bim"],
        job_count=2,
        total_job_count=2,
        unique_job_count=2,
        relevant_job_count=1,
        job_titles=["BIM Coordinator"],
        locations=["Oslo, Norway"],
        technologies=["Revit", "Navisworks"],
        signals=[{"signal": "bim_coordination", "evidence_count": 1}],
        evidence=["BIM coordination workflow"],
        fit_score=20,       # <= 30 ✓
        intent_score=10,    # <= 30 ✓
        recency_score=20,   # <= 20 ✓
        evidence_score=10,  # <= 20 ✓
        lead_score=60,      # 20 + 10 + 20 + 10 = 60 ✓
        qualification_threshold=60,
        qualified=True,
        scoring_reasons=["BIM tech detected"],
    )
    c2_1 = ContactCandidate(
        first_name="Erik",
        last_name="Hansen",
        full_name="Erik Hansen",
        job_title="BIM Manager",
        seniority="Manager",
        department="Digital Delivery",
        company_name="Nordic BIM Consulting AS",
        company_domain="nordicbim.no",
        work_email="erik.hansen@nordicbim.no",
        email_status="likely",
        email_confidence=82,
        buyer_role_match="exact",
        contact_score=80,
        data_sources=["apollo"],
        is_best_contact=True,
    )
    e2 = EnrichedLead(
        base_lead=l2,
        company_enrichment=CompanyEnrichment(
            company_name="Nordic BIM Consulting AS",
            company_domain="nordicbim.no",
            industry="Engineering Services",
            employee_count=45,
            country="Norway",
        ),
        contacts=[c2_1],
        best_contact=c2_1,
        enrichment_status="partial",
        providers_used=["apollo"],
        enrichment_errors=[],
    )

    return [e1, e2]


def build_sample_jobs():
    return [
        StructuredJob(
            company_name="VDC Automation Lab GmbH",
            company_domain="vdc-autolab.de",
            source="lever",
            source_company_key="vdc-autolab",
            job_title="Head of Digital Practice",
            job_url="https://jobs.lever.co/vdc-autolab/101",
            location="Munich, Germany",
            technologies=["Revit", "Dynamo", "Python"],
            relevant_signals=[],
        ),
        StructuredJob(
            company_name="VDC Automation Lab GmbH",
            company_domain="vdc-autolab.de",
            source="lever",
            source_company_key="vdc-autolab",
            job_title="Revit API Developer",
            job_url="https://jobs.lever.co/vdc-autolab/102",
            location="Remote",
            technologies=["Revit", "C#"],
            relevant_signals=[],
        ),
        StructuredJob(
            company_name="Nordic BIM Consulting AS",
            company_domain="nordicbim.no",
            source="greenhouse",
            source_company_key="nordic-bim",
            job_title="BIM Coordinator",
            job_url="https://boards.greenhouse.io/nordicbim/301",
            location="Oslo, Norway",
            technologies=["Revit", "Navisworks"],
            relevant_signals=[],
        ),
    ]


def main():
    # Clear sample directory
    if os.path.exists(SAMPLE_DIR):
        shutil.rmtree(SAMPLE_DIR)
    os.makedirs(SAMPLE_DIR, exist_ok=True)

    leads = build_sample_leads()
    jobs = build_sample_jobs()

    # Validate Phase 5 scoring contract before export
    for enriched in leads:
        base = enriched.base_lead
        assert base.fit_score <= 30, f"{base.company_name}: fit_score={base.fit_score} > 30"
        assert base.intent_score <= 30, f"{base.company_name}: intent_score={base.intent_score} > 30"
        assert base.recency_score <= 20, f"{base.company_name}: recency_score={base.recency_score} > 20"
        assert base.evidence_score <= 20, f"{base.company_name}: evidence_score={base.evidence_score} > 20"
        component_sum = base.fit_score + base.intent_score + base.recency_score + base.evidence_score
        assert base.lead_score == component_sum, (
            f"{base.company_name}: lead_score={base.lead_score} != component sum {component_sum}"
        )
        assert base.qualification_threshold == 60, f"{base.company_name}: threshold={base.qualification_threshold}"
        assert base.qualified == (base.lead_score >= base.qualification_threshold)

    print("✓ Phase 5 scoring contract validated for all sample leads")

    orchestrator = ExportOrchestrator()
    req = ExportRequest(
        leads=leads,
        jobs=jobs,
        formats=["xlsx", "csv", "json"],
        qualified_only=False,
        include_all_contacts=True,
        include_jobs=True,
        output_dir=SAMPLE_DIR,
    )

    resp = orchestrator.export(req)

    print(f"\n✓ Sample export generated successfully!")
    print(f"  Export Run ID: {resp.export_run_id}")
    print(f"  Leads exported: {resp.exported_leads}")
    print(f"  Contacts exported: {resp.exported_contacts}")
    print(f"  Jobs exported: {resp.exported_jobs}")
    print(f"  Errors: {len(resp.errors)}")
    print(f"\n  Output directory: {SAMPLE_DIR}")
    for fmt, path in resp.files.items():
        print(f"  {fmt}: {path}")

    # Post-export validation
    import openpyxl
    xlsx_path = resp.files.get("xlsx")
    if xlsx_path:
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb["Leads"]
        for row_idx in range(2, ws.max_row + 1):
            fit = ws.cell(row=row_idx, column=10).value or 0
            intent = ws.cell(row=row_idx, column=11).value or 0
            recency = ws.cell(row=row_idx, column=12).value or 0
            evidence = ws.cell(row=row_idx, column=13).value or 0
            total = ws.cell(row=row_idx, column=8).value or 0
            threshold = ws.cell(row=row_idx, column=16).value or 60
            name = ws.cell(row=row_idx, column=3).value
            
            assert fit <= 30, f"Row {row_idx} ({name}): fit={fit} > 30"
            assert intent <= 30, f"Row {row_idx} ({name}): intent={intent} > 30"
            assert recency <= 20, f"Row {row_idx} ({name}): recency={recency} > 20"
            assert evidence <= 20, f"Row {row_idx} ({name}): evidence={evidence} > 20"
            assert total == fit + intent + recency + evidence, f"Row {row_idx} ({name}): score mismatch"
            assert threshold == 60, f"Row {row_idx} ({name}): threshold={threshold}, expected 60"

        print(f"\n✓ Post-export XLSX validation passed (all rows comply with Phase 5 contract)")
    
    print("\n✓ Sample regeneration complete. Ready for Phase 7 freeze.")


if __name__ == "__main__":
    main()
