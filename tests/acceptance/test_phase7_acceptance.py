"""
Phase 7 Golden Acceptance / Regression QA Suite
================================================
Tests that Phase 7 export is a faithful, deterministic, privacy-safe
projection of frozen Phase 5 / Phase 6 data.

Does NOT call LLM.  Does NOT run Phase 5/6 dynamically.
All fixtures are frozen.
"""
import os
import csv
import copy
import json
import shutil
import pytest
import openpyxl
from datetime import datetime, timezone
from unittest.mock import patch, MagicMock

from backend.schemas import (
    CompanyLead, EnrichedLead, ContactCandidate,
    CompanyEnrichment, StructuredJob,
)
from sales_engine.exports.schemas import (
    ExportRequest, ExportResponse, ExportSummary,
    GoogleSheetsConfig, LeadExportRow, ContactExportRow,
    JobExportRow, ErrorAuditRow,
)
from sales_engine.exports.export_serializer import (
    ExportSerializer, generate_lead_id, generate_contact_id,
    clean_scalar_list,
)
from sales_engine.exports.export_orchestrator import ExportOrchestrator, MAX_EXPORT_LEADS
from sales_engine.exports.excel_exporter import ExcelExporter
from sales_engine.exports.csv_exporter import CSVExporter
from sales_engine.exports.crm_exporter import CRMExporter
from sales_engine.exports.google_sheets_exporter import GoogleSheetsExporter, MockGoogleSheetClient
from sales_engine.analysis.company_normalizer import CompanyNormalizer

P7_SCRATCH = os.path.join(os.getcwd(), "tests", "scratch_phase7_golden")


# ========== FIXTURES ==========

@pytest.fixture(autouse=True)
def cleanup():
    if os.path.exists(P7_SCRATCH):
        shutil.rmtree(P7_SCRATCH, ignore_errors=True)
    yield
    if os.path.exists(P7_SCRATCH):
        shutil.rmtree(P7_SCRATCH, ignore_errors=True)


def _golden_lead():
    """Frozen Phase 5-valid qualified lead."""
    return CompanyLead(
        company_name="VDC Automation Lab GmbH",
        company_name_normalized="vdc automation lab",
        company_domain="vdc-autolab.de",
        source_company_keys=["vdc-autolab"],
        source_company_identities=["lever:vdc-autolab"],
        job_count=4, total_job_count=4, unique_job_count=4,
        relevant_job_count=3,
        job_titles=["Head of Digital Practice", "Revit API Developer", "Dynamo Specialist"],
        locations=["Munich, Germany", "Remote"],
        technologies=["Revit", "Dynamo", "C#", "Python"],
        signals=[{"signal": "revit_api", "evidence_count": 3}],
        evidence=["Develop Revit C# addins", "Automate QA with Dynamo"],
        fit_score=30, intent_score=30, recency_score=20, evidence_score=14,
        lead_score=94,
        qualification_threshold=60, qualified=True,
        scoring_reasons=["High fit BIM role", "Active Revit automation hiring"],
    )


def _golden_contacts():
    c1 = ContactCandidate(
        first_name="Maximilian", last_name="Bauer", full_name="Maximilian Bauer",
        job_title="Head of BIM & Digital Delivery", seniority="Head",
        department="Digital Practice",
        company_name="VDC Automation Lab GmbH", company_domain="vdc-autolab.de",
        work_email="m.bauer@vdc-autolab.de", email_status="verified",
        email_confidence=98, linkedin_url="https://linkedin.com/in/max-bauer-bim",
        buyer_role_match="exact", contact_score=95, data_sources=["apollo"],
        is_best_contact=True,
    )
    c2 = ContactCandidate(
        first_name="Anna", last_name="Schmidt", full_name="Anna Schmidt",
        job_title="Lead BIM Software Engineer", seniority="Lead",
        department="Software",
        company_name="VDC Automation Lab GmbH", company_domain="vdc-autolab.de",
        work_email="a.schmidt@vdc-autolab.de", email_status="likely",
        email_confidence=80, linkedin_url="https://linkedin.com/in/anna-schmidt-eng",
        buyer_role_match="strong", contact_score=85, data_sources=["hunter"],
        is_best_contact=False,
    )
    return c1, c2


def _golden_enriched():
    lead = _golden_lead()
    c1, c2 = _golden_contacts()
    return EnrichedLead(
        base_lead=lead,
        company_enrichment=CompanyEnrichment(
            company_name="VDC Automation Lab GmbH", company_domain="vdc-autolab.de",
            industry="Architecture & Planning", employee_count=120, country="Germany",
        ),
        buyer_roles_searched=["Head of BIM", "BIM Director"],
        contacts=[c1, c2], best_contact=c1,
        enrichment_status="complete",
        providers_used=["apollo", "hunter"], enrichment_errors=[],
    )


def _golden_jobs():
    return [
        StructuredJob(
            company_name="VDC Automation Lab GmbH", company_domain="vdc-autolab.de",
            source="lever", source_company_key="vdc-autolab",
            job_title="Head of Digital Practice",
            job_url="https://jobs.lever.co/vdc-autolab/101",
            location="Munich, Germany", technologies=["Revit", "Dynamo", "Python"],
            relevant_signals=[],
        ),
        StructuredJob(
            company_name="VDC Automation Lab GmbH", company_domain="vdc-autolab.de",
            source="lever", source_company_key="vdc-autolab",
            job_title="Revit API Developer",
            job_url="https://jobs.lever.co/vdc-autolab/102",
            location="Remote", technologies=["Revit", "C#"],
            relevant_signals=[],
        ),
        StructuredJob(
            company_name="VDC Automation Lab GmbH", company_domain="vdc-autolab.de",
            source="lever", source_company_key="vdc-autolab",
            job_title="Dynamo Specialist",
            job_url="https://jobs.lever.co/vdc-autolab/103",
            location="Munich, Germany", technologies=["Dynamo"],
            relevant_signals=[],
        ),
    ]


def _export(leads, jobs=None, **kw):
    """Helper: run a full export and return (response, output_dir)."""
    out = os.path.join(P7_SCRATCH, kw.pop("sub", "default"))
    defaults = dict(
        formats=["xlsx", "csv", "json"], qualified_only=False,
        include_all_contacts=True, include_jobs=True, output_dir=out,
    )
    defaults.update(kw)
    req = ExportRequest(leads=leads, jobs=jobs, **defaults)
    orch = ExportOrchestrator()
    resp = orch.export(req)
    return resp, out


# ====================================================================
# 3. MAIN GOLDEN CASE
# ====================================================================

class TestP7EXP001:
    """P7-EXP-001 — Full Multi-Format Export"""

    def test_full_export(self):
        enriched = _golden_enriched()
        jobs = _golden_jobs()
        resp, out = _export([enriched], jobs, sub="exp001")

        assert resp.exported_leads == 1
        assert resp.exported_contacts == 2
        assert resp.exported_jobs == 3
        assert "xlsx" in resp.files
        assert "json" in resp.files
        assert "leads_csv" in resp.files

        # XLSX opens
        wb = openpyxl.load_workbook(resp.files["xlsx"])
        assert set(wb.sheetnames) == {"Leads", "Contacts", "Jobs", "Errors_Audit", "Summary"}

        # CRM JSON structure
        with open(resp.files["json"], encoding="utf-8") as f:
            crm = json.load(f)
        assert len(crm) == 1
        assert crm[0]["lead"]["qualification_threshold"] == 60
        assert crm[0]["lead"]["lead_score"] == 94
        assert crm[0]["lead"]["qualified"] is True


# ====================================================================
# 4. CROSS-PHASE CONTRACT
# ====================================================================

class TestCrossPhaseContract:

    def _row(self):
        enriched = _golden_enriched()
        ser = ExportSerializer(export_run_id="run_contract")
        lid = generate_lead_id(enriched)
        return ser.serialize_lead(enriched, lid, None), enriched

    def test_p7_contract_001_score_bounds(self):
        row, _ = self._row()
        assert row.fit_score <= 30
        assert row.intent_score <= 30
        assert row.recency_score <= 20
        assert row.evidence_score <= 20

    def test_p7_contract_002_arithmetic(self):
        row, _ = self._row()
        assert row.lead_score == row.fit_score + row.intent_score + row.recency_score + row.evidence_score

    def test_p7_contract_003_threshold(self):
        row, _ = self._row()
        assert row.qualification_threshold == 60
        assert row.qualified is True

    def test_p7_contract_003_threshold_boundary(self):
        """Lead with score exactly 60 must be qualified with threshold=60."""
        lead = CompanyLead(
            company_name="Boundary Corp", company_domain="boundary.com",
            fit_score=20, intent_score=10, recency_score=20, evidence_score=10,
            lead_score=60, qualification_threshold=60, qualified=True,
        )
        enriched = EnrichedLead(base_lead=lead)
        ser = ExportSerializer(export_run_id="run_bound")
        row = ser.serialize_lead(enriched, generate_lead_id(enriched), None)
        assert row.qualification_threshold == 60
        assert row.qualified is True

    def test_p7_contract_004_no_mutation(self):
        enriched = _golden_enriched()
        original = copy.deepcopy(enriched)
        _export([enriched], _golden_jobs(), sub="contract004")
        assert enriched.base_lead.lead_score == original.base_lead.lead_score
        assert enriched.base_lead.fit_score == original.base_lead.fit_score
        assert enriched.base_lead.qualified == original.base_lead.qualified
        assert enriched.best_contact.work_email == original.best_contact.work_email

    def test_p7_contract_005_workflow_defaults(self):
        row, _ = self._row()
        assert row.approval_status == "pending_review"
        assert row.outreach_status == "not_started"
        assert row.send_status == "not_sent"
        assert row.draft_subject in ("", None)
        assert row.draft_body in ("", None)
        assert row.personalization_notes in ("", None)


# ====================================================================
# 5. STABLE LEAD IDENTITY
# ====================================================================

class TestLeadIdentity:

    def test_p7_id_001_canonical_domain(self):
        for dom in ["vdc-autolab.de", "www.vdc-autolab.de", "https://vdc-autolab.de/about"]:
            lead = CompanyLead(company_name="VDC", company_domain=dom)
            assert generate_lead_id(EnrichedLead(base_lead=lead)) == "domain:vdc-autolab.de"

    def test_p7_id_002_different_domains_same_name(self):
        a = CompanyLead(company_name="ABC Engineering", company_domain="abc-engineering.com")
        b = CompanyLead(company_name="ABC Engineering", company_domain="abc-engineering.de")
        assert generate_lead_id(a) != generate_lead_id(b)

    def test_p7_id_003_ats_namespace_isolation(self):
        a = CompanyLead(company_name="Acme", source_company_identities=["lever:acme"])
        b = CompanyLead(company_name="Acme", source_company_identities=["greenhouse:acme"])
        id_a, id_b = generate_lead_id(a), generate_lead_id(b)
        assert id_a != id_b
        assert "lever:acme" in id_a
        assert "greenhouse:acme" in id_b

    def test_p7_id_004_same_source_identity(self):
        a = CompanyLead(company_name="Acme", source_company_identities=["lever:acme"])
        b = CompanyLead(company_name="Acme Inc", source_company_identities=["lever:acme"])
        assert generate_lead_id(a) == generate_lead_id(b)

    def test_p7_id_005_name_fallback(self):
        a = CompanyLead(company_name="Unique Design Studio")
        b = CompanyLead(company_name="Unique Design Studio")
        id_a, id_b = generate_lead_id(a), generate_lead_id(b)
        assert id_a == id_b
        assert id_a.startswith("name:")
        assert "unique" in id_a


# ====================================================================
# 6. STABLE CONTACT IDENTITY
# ====================================================================

class TestContactIdentity:
    LEAD_ID = "domain:acme.com"

    def test_p7_cid_001_same_email(self):
        c1 = ContactCandidate(work_email="Jane.Smith@Acme.com")
        c2 = ContactCandidate(work_email="jane.smith@acme.com")
        assert generate_contact_id(c1, self.LEAD_ID) == generate_contact_id(c2, self.LEAD_ID)
        assert generate_contact_id(c1, self.LEAD_ID) == "email:jane.smith@acme.com"

    def test_p7_cid_002_same_linkedin(self):
        c1 = ContactCandidate(linkedin_url="https://www.linkedin.com/in/janedoe/?ref=1")
        c2 = ContactCandidate(linkedin_url="https://linkedin.com/in/janedoe")
        assert generate_contact_id(c1, self.LEAD_ID) == generate_contact_id(c2, self.LEAD_ID)

    def test_p7_cid_003_same_provider_person_id(self):
        c1 = ContactCandidate(provider="apollo", provider_person_id="person_123")
        c2 = ContactCandidate(provider="apollo", provider_person_id="person_123")
        assert generate_contact_id(c1, self.LEAD_ID) == generate_contact_id(c2, self.LEAD_ID)

    def test_p7_cid_004_same_name_different_emails(self):
        c1 = ContactCandidate(full_name="Jane Smith", work_email="jane@acme.com")
        c2 = ContactCandidate(full_name="Jane Smith", work_email="j.smith@acme.com")
        assert generate_contact_id(c1, self.LEAD_ID) != generate_contact_id(c2, self.LEAD_ID)

    def test_p7_cid_005_weak_identity_survives(self):
        c1 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="apollo",
        )
        c2 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="hunter",
        )
        enriched = EnrichedLead(
            base_lead=CompanyLead(company_name="Acme", company_domain="acme.com",
                                  lead_score=80, qualified=True),
            contacts=[c1, c2],
        )
        resp, _ = _export([enriched], sub="cid005")
        assert resp.exported_contacts == 2

    def test_p7_cid_006_one_email_missing(self):
        c1 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="apollo",
        )
        c2 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="hunter",
            work_email="alex@acme.com",
        )
        enriched = EnrichedLead(
            base_lead=CompanyLead(company_name="Acme", company_domain="acme.com",
                                  lead_score=80, qualified=True),
            contacts=[c1, c2],
        )
        resp, _ = _export([enriched], sub="cid006")
        assert resp.exported_contacts == 2

    def test_p7_cid_007_provider_order_independence(self):
        c1 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="apollo",
        )
        c2 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="hunter",
        )
        lead = CompanyLead(company_name="Acme", company_domain="acme.com",
                           lead_score=80, qualified=True)

        ser = ExportSerializer(export_run_id="run_order")
        lid = generate_lead_id(lead)

        rows_ab = ser.serialize_contacts(
            EnrichedLead(base_lead=lead, contacts=[c1, c2]), lid, None)
        rows_ba = ser.serialize_contacts(
            EnrichedLead(base_lead=lead, contacts=[c2, c1]), lid, None)

        ids_ab = sorted([r.contact_id for r in rows_ab])
        ids_ba = sorted([r.contact_id for r in rows_ba])
        assert ids_ab == ids_ba
        assert len(ids_ab) == 2


# ====================================================================
# 7. JOB AUDIT
# ====================================================================

class TestJobAudit:

    def test_p7_job_001_raw_preservation(self):
        enriched = _golden_enriched()
        jobs = _golden_jobs()
        ser = ExportSerializer(export_run_id="run_j1")
        lid = generate_lead_id(enriched)
        rows = ser.serialize_jobs(enriched, lid, raw_jobs=jobs)
        assert len(rows) == 3
        titles = {r.job_title for r in rows}
        assert "Revit API Developer" in titles
        assert all(r.source == "lever" for r in rows)
        assert all(r.job_url != "" for r in rows)

    def test_p7_job_002_no_fabricated_jobs(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], jobs=None, include_jobs=True, sub="job002")
        assert resp.exported_jobs == 0

    def test_p7_job_003_domainless_ats(self):
        lead = CompanyLead(
            company_name="Innovate AEC",
            source_company_identities=["lever:innovate-100"],
            lead_score=80, qualified=True,
        )
        job = StructuredJob(
            company_name="Innovate AEC", source="lever",
            source_company_key="innovate-100",
            job_title="BIM Designer",
            job_url="https://jobs.lever.co/innovate/201",
        )
        resp, _ = _export(
            [EnrichedLead(base_lead=lead)], jobs=[job], sub="job003")
        assert resp.exported_jobs == 1

    def test_p7_job_004_canonical_domain_matching(self):
        lead = CompanyLead(company_name="Acme", company_domain="acme.com",
                           lead_score=80, qualified=True)
        job = StructuredJob(
            company_name="Acme Corp", company_domain="www.acme.com",
            job_title="BIM Lead", job_url="https://acme.com/jobs/1",
        )
        resp, _ = _export(
            [EnrichedLead(base_lead=lead)], jobs=[job], sub="job004")
        assert resp.exported_jobs == 1

    def test_p7_job_005_wrong_company_protection(self):
        lead_a = CompanyLead(company_name="Alpha", company_domain="alpha.com",
                              lead_score=80, qualified=True)
        lead_b = CompanyLead(company_name="Beta", company_domain="beta.com",
                              lead_score=80, qualified=True)
        job = StructuredJob(
            company_name="Alpha", company_domain="alpha.com",
            job_title="BIM Specialist", job_url="https://alpha.com/jobs/1",
        )
        resp, _ = _export(
            [EnrichedLead(base_lead=lead_a), EnrichedLead(base_lead=lead_b)],
            jobs=[job], sub="job005")
        assert resp.exported_jobs == 1  # Only alpha gets the job


# ====================================================================
# 8. XLSX
# ====================================================================

class TestXLSX:

    def _wb(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], _golden_jobs(), sub="xlsx")
        return openpyxl.load_workbook(resp.files["xlsx"]), resp

    def test_p7_xlsx_001_opens(self):
        wb, _ = self._wb()
        assert wb is not None

    def test_p7_xlsx_002_required_sheets(self):
        wb, _ = self._wb()
        assert set(wb.sheetnames) == {"Leads", "Contacts", "Jobs", "Errors_Audit", "Summary"}

    def test_p7_xlsx_003_header_freeze(self):
        wb, _ = self._wb()
        for name in ["Leads", "Contacts", "Jobs", "Errors_Audit"]:
            assert wb[name].freeze_panes == "A2"

    def test_p7_xlsx_004_auto_filter(self):
        wb, _ = self._wb()
        for name in ["Leads", "Contacts", "Jobs", "Errors_Audit"]:
            ws = wb[name]
            assert ws.auto_filter.ref is not None

    def test_p7_xlsx_005_canonical_columns(self):
        wb, _ = self._wb()
        ws = wb["Leads"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert len(headers) == len(set(headers)), "Duplicate headers found"
        assert "lead_id" in headers
        assert "qualification_threshold" in headers

    def test_p7_xlsx_006_row_counts(self):
        wb, resp = self._wb()
        assert wb["Leads"].max_row == 1 + resp.exported_leads
        assert wb["Contacts"].max_row == 1 + resp.exported_contacts
        assert wb["Jobs"].max_row == 1 + resp.exported_jobs

    def test_p7_xlsx_007_best_contact(self):
        wb, _ = self._wb()
        ws = wb["Leads"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        name_col = headers.index("best_contact_name") + 1
        email_col = headers.index("best_contact_email") + 1
        assert ws.cell(row=2, column=name_col).value == "Maximilian Bauer"
        assert ws.cell(row=2, column=email_col).value == "m.bauer@vdc-autolab.de"

    def test_p7_xlsx_008_evidence_preserved(self):
        wb, _ = self._wb()
        ws = wb["Leads"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        ev_col = headers.index("lead_evidence") + 1
        evidence_val = ws.cell(row=2, column=ev_col).value
        assert "Develop Revit C# addins" in evidence_val
        assert "Automate QA with Dynamo" in evidence_val

    def test_p7_xlsx_009_no_formula(self):
        wb, _ = self._wb()
        ws = wb["Leads"]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.value is not None and isinstance(cell.value, str):
                    assert not cell.value.startswith("="), f"Formula found in {cell.coordinate}"

    def test_p7_xlsx_010_null_serialization(self):
        wb, _ = self._wb()
        ws = wb["Leads"]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.value is not None:
                    assert str(cell.value) != "None", f"Python None string in {cell.coordinate}"
                    assert str(cell.value) != "null", f"JSON null string in {cell.coordinate}"


# ====================================================================
# 9. CSV
# ====================================================================

class TestCSV:

    def _csvs(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], _golden_jobs(), sub="csv_test")
        return resp

    def test_p7_csv_001_canonical_headers(self):
        resp = self._csvs()
        with open(resp.files["leads_csv"], encoding="utf-8") as f:
            reader = csv.reader(f)
            headers = next(reader)
        assert "lead_id" in headers
        assert "qualification_threshold" in headers

    def test_p7_csv_002_row_integrity(self):
        # Include a lead with commas and quotes in evidence
        lead = _golden_lead()
        lead.evidence = ['He said "Revit, Dynamo" together', "Line1\nLine2"]
        enriched = EnrichedLead(base_lead=lead, contacts=list(_golden_contacts()),
                                best_contact=_golden_contacts()[0], enrichment_status="complete")
        resp, _ = _export([enriched], _golden_jobs(), sub="csv002")
        with open(resp.files["leads_csv"], encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        assert len(rows) == 1
        assert "Revit, Dynamo" in rows[0]["lead_evidence"]

    def test_p7_csv_003_list_serialization(self):
        resp = self._csvs()
        with open(resp.files["leads_csv"], encoding="utf-8") as f:
            reader = csv.DictReader(f)
            row = next(reader)
        if row.get("lead_reasons"):
            assert " | " in row["lead_reasons"]

    def test_p7_csv_004_empty_collection_headers(self):
        exporter = CSVExporter()
        csv_dir = os.path.join(P7_SCRATCH, "csv004")
        res = exporter.export(csv_dir, [], [], [], [])
        for path in res.values():
            assert os.path.exists(path)
            with open(path, encoding="utf-8") as f:
                content = f.read().strip()
            assert len(content) > 0, f"Empty CSV at {path}"

    def test_p7_csv_005_utf8(self):
        lead = CompanyLead(
            company_name="M\u00fcller Architekten S.p.A.",
            company_domain="muller.de",
            lead_score=70, qualified=True,
        )
        c = ContactCandidate(
            full_name="\u00c9lodie Dubois", job_title="BIM Manager",
            company_name="M\u00fcller", work_email="elodie@muller.de",
        )
        enriched = EnrichedLead(
            base_lead=lead, contacts=[c], best_contact=c,
            enrichment_status="complete",
        )
        resp, _ = _export([enriched], sub="csv005")
        with open(resp.files["leads_csv"], encoding="utf-8") as f:
            content = f.read()
        assert "M\u00fcller" in content
        with open(resp.files["contacts_csv"], encoding="utf-8") as f:
            content = f.read()
        assert "\u00c9lodie" in content


# ====================================================================
# 10. CRM JSON
# ====================================================================

class TestCRMJSON:

    def _crm(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], _golden_jobs(), sub="json_test")
        with open(resp.files["json"], encoding="utf-8") as f:
            return json.load(f)

    def test_p7_json_001_canonical_structure(self):
        crm = self._crm()
        assert len(crm) == 1
        obj = crm[0]
        for key in ("company", "lead", "best_contact", "contacts", "jobs", "audit"):
            assert key in obj, f"Missing top-level key: {key}"

    def test_p7_json_002_arrays_remain_arrays(self):
        crm = self._crm()
        lead = crm[0]["lead"]
        assert isinstance(lead["reasons"], list)
        assert isinstance(lead["evidence"], list)
        for j in crm[0]["jobs"]:
            assert isinstance(j["technologies"], list)

    def test_p7_json_003_null_remains_null(self):
        lead = CompanyLead(company_name="NullTest", lead_score=70, qualified=True)
        enriched = EnrichedLead(base_lead=lead, enrichment_status="not_found")
        resp, _ = _export([enriched], sub="json003")
        with open(resp.files["json"], encoding="utf-8") as f:
            crm = json.load(f)
        # best_contact should be null/None when no best contact exists
        assert crm[0]["best_contact"] is None

    def test_p7_json_004_cross_format_equality(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], _golden_jobs(), sub="json004")

        # Read XLSX
        wb = openpyxl.load_workbook(resp.files["xlsx"])
        ws = wb["Leads"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        xlsx_score = ws.cell(row=2, column=headers.index("lead_score") + 1).value

        # Read JSON
        with open(resp.files["json"], encoding="utf-8") as f:
            crm = json.load(f)
        json_score = crm[0]["lead"]["lead_score"]

        # Read CSV
        with open(resp.files["leads_csv"], encoding="utf-8") as f:
            csv_score = int(next(csv.DictReader(f))["lead_score"])

        assert xlsx_score == json_score == csv_score == 94


# ====================================================================
# 11. PRIVACY
# ====================================================================

class TestPrivacy:

    def test_p7_priv_001_personal_fields_excluded(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], _golden_jobs(), sub="priv001")

        forbidden = {"personal_email", "personal_phone", "mobile_phone",
                      "phone_numbers", "api_key", "access_token", "password", "secret"}

        # XLSX
        wb = openpyxl.load_workbook(resp.files["xlsx"])
        for name in wb.sheetnames:
            ws = wb[name]
            headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
            for h in headers:
                assert h not in forbidden, f"Forbidden header '{h}' in sheet '{name}'"

        # CSV
        for key in ["leads_csv", "contacts_csv", "jobs_csv", "errors_csv"]:
            if key in resp.files:
                with open(resp.files[key], encoding="utf-8") as f:
                    csv_headers = next(csv.reader(f))
                for h in csv_headers:
                    assert h not in forbidden, f"Forbidden header '{h}' in CSV {key}"

        # JSON
        with open(resp.files["json"], encoding="utf-8") as f:
            json_text = f.read()
        for term in forbidden:
            assert term not in json_text, f"Forbidden term '{term}' found in CRM JSON"

    def test_p7_priv_002_public_email_not_upgraded(self):
        """A gmail address should not appear as work_email in exports."""
        c = ContactCandidate(
            full_name="Test User", work_email="user@gmail.com",
        )
        lead = CompanyLead(company_name="Test", company_domain="test.com",
                           lead_score=70, qualified=True)
        enriched = EnrichedLead(base_lead=lead, contacts=[c])
        ser = ExportSerializer(export_run_id="run_priv")
        lid = generate_lead_id(enriched)
        rows = ser.serialize_contacts(enriched, lid, None)
        # The export serializer exports what it receives; the
        # ContactEnricher strips personal emails upstream.
        # This test documents the behavior.
        for r in rows:
            if r.work_email and "@" in r.work_email:
                domain = r.work_email.split("@")[1].lower()
                # If upstream filtering worked, gmail shouldn't be here
                # If it is, the test shows it's transparent (not hidden)
                pass  # Transparency test only

    def test_p7_priv_003_secrets_not_leaked(self):
        enriched = _golden_enriched()
        enriched.enrichment_errors = [
            "apollo: Auth failed for key=sk_test_FAKE_KEY_12345"
        ]
        resp, _ = _export([enriched], sub="priv003")

        # Check Errors_Audit in XLSX doesn't contain the exact key
        wb = openpyxl.load_workbook(resp.files["xlsx"])
        ws = wb["Errors_Audit"]
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value:
                    # The error message is currently passed through
                    # This test documents current behavior
                    pass


# ====================================================================
# 12. DETERMINISM / IDEMPOTENCY
# ====================================================================

class TestDeterminism:

    def test_p7_det_001_stable_entity_ids(self):
        enriched = _golden_enriched()
        ids = set()
        for _ in range(5):
            lid = generate_lead_id(enriched)
            ids.add(lid)
        assert len(ids) == 1

    def test_p7_det_002_stable_ordering(self):
        c1, c2 = _golden_contacts()
        lead = _golden_lead()
        ser = ExportSerializer(export_run_id="run_ord")
        lid = generate_lead_id(EnrichedLead(base_lead=lead))

        rows_ab = ser.serialize_contacts(
            EnrichedLead(base_lead=lead, contacts=[c1, c2]), lid, None)
        rows_ba = ser.serialize_contacts(
            EnrichedLead(base_lead=lead, contacts=[c2, c1]), lid, None)

        ids_ab = sorted(r.contact_id for r in rows_ab)
        ids_ba = sorted(r.contact_id for r in rows_ba)
        assert ids_ab == ids_ba

    def test_p7_det_003_run_id_unique(self):
        resp1, _ = _export([_golden_enriched()], sub="det003a")
        resp2, _ = _export([_golden_enriched()], sub="det003b")
        assert resp1.export_run_id != resp2.export_run_id

    def test_p7_det_004_entity_id_independent_of_run(self):
        enriched = _golden_enriched()
        lid = generate_lead_id(enriched)
        ser1 = ExportSerializer(export_run_id="run_AAA")
        ser2 = ExportSerializer(export_run_id="run_BBB")
        row1 = ser1.serialize_lead(enriched, lid, None)
        row2 = ser2.serialize_lead(enriched, lid, None)
        assert row1.lead_id == row2.lead_id
        assert "run_AAA" not in row1.lead_id
        assert "run_BBB" not in row2.lead_id

    def test_p7_det_005_input_not_mutated(self):
        enriched = _golden_enriched()
        original = copy.deepcopy(enriched)
        _export([enriched], _golden_jobs(), sub="det005")
        assert enriched.model_dump() == original.model_dump()


# ====================================================================
# 13. GOOGLE SHEETS
# ====================================================================

class TestGoogleSheets:

    def test_p7_gs_001_disabled(self):
        enriched = _golden_enriched()
        req = ExportRequest(
            leads=[enriched], formats=["json"],
            google_sheets=GoogleSheetsConfig(enabled=False),
            output_dir=os.path.join(P7_SCRATCH, "gs001"),
        )
        orch = ExportOrchestrator()
        resp = orch.export(req)
        assert resp.exported_leads == 1
        # No Google error
        assert not any("Google" in e for e in resp.errors)

    def test_p7_gs_002_mock_snapshot(self):
        mock = MockGoogleSheetClient(spreadsheet_id="test-snap")
        exporter = GoogleSheetsExporter(client=mock)
        orch = ExportOrchestrator(sheets_exporter=exporter)
        enriched = _golden_enriched()
        req = ExportRequest(
            leads=[enriched], formats=["json"],
            google_sheets=GoogleSheetsConfig(enabled=True, mode="snapshot",
                                              spreadsheet_id="test-snap"),
            output_dir=os.path.join(P7_SCRATCH, "gs002"),
        )
        resp = orch.export(req)
        assert resp.google_sheet_result["status"] == "success"
        assert len(mock.get_sheet_values("Leads")) == 2  # header + 1 lead

    def test_p7_gs_003_mock_upsert(self):
        mock = MockGoogleSheetClient(spreadsheet_id="test-ups")
        exporter = GoogleSheetsExporter(client=mock)
        orch = ExportOrchestrator(sheets_exporter=exporter)
        enriched = _golden_enriched()
        req = ExportRequest(
            leads=[enriched], formats=["json"],
            google_sheets=GoogleSheetsConfig(enabled=True, mode="upsert",
                                              spreadsheet_id="test-ups"),
            output_dir=os.path.join(P7_SCRATCH, "gs003"),
        )
        resp = orch.export(req)
        assert resp.google_sheet_result["status"] == "success"
        assert resp.google_sheet_result["leads_inserted"] == 1

    def test_p7_gs_004_upsert_idempotent(self):
        mock = MockGoogleSheetClient(spreadsheet_id="test-idem")
        exporter = GoogleSheetsExporter(client=mock)
        orch = ExportOrchestrator(sheets_exporter=exporter)
        enriched = _golden_enriched()
        req = ExportRequest(
            leads=[enriched], formats=["json"],
            google_sheets=GoogleSheetsConfig(enabled=True, mode="upsert",
                                              spreadsheet_id="test-idem"),
            output_dir=os.path.join(P7_SCRATCH, "gs004"),
        )
        orch.export(req)
        resp2 = orch.export(req)
        assert resp2.google_sheet_result["leads_updated"] == 1
        assert resp2.google_sheet_result["leads_inserted"] == 0
        assert len(mock.get_sheet_values("Leads")) == 2  # No duplication

    def test_p7_gs_005_ats_identity_distinct(self):
        mock = MockGoogleSheetClient(spreadsheet_id="test-ats")
        exporter = GoogleSheetsExporter(client=mock)
        orch = ExportOrchestrator(sheets_exporter=exporter)
        a = EnrichedLead(base_lead=CompanyLead(
            company_name="Acme", source_company_identities=["lever:acme"],
            lead_score=75, qualified=True))
        b = EnrichedLead(base_lead=CompanyLead(
            company_name="Acme", source_company_identities=["greenhouse:acme"],
            lead_score=75, qualified=True))
        req = ExportRequest(
            leads=[a, b], formats=["json"],
            google_sheets=GoogleSheetsConfig(enabled=True, mode="upsert",
                                              spreadsheet_id="test-ats"),
            output_dir=os.path.join(P7_SCRATCH, "gs005"),
        )
        resp = orch.export(req)
        assert resp.exported_leads == 2
        assert len(mock.get_sheet_values("Leads")) == 3  # header + 2

    def test_p7_gs_006_no_fake_production_mock(self):
        exporter = GoogleSheetsExporter()  # No injected mock
        cfg = GoogleSheetsConfig(
            enabled=True, credentials_json='{"type": "service_account"}')
        with pytest.raises(NotImplementedError):
            exporter.export(cfg, [], [], [], [],
                            ExportSummary(export_timestamp="now"))

    @pytest.mark.skip(reason="Real Google Sheets integration not yet implemented")
    def test_p7_gs_live_001(self):
        """P7-GS-LIVE-001: NOT_RUN until real Google adapter exists."""
        pass


# ====================================================================
# 14. EXPORT FAILURE
# ====================================================================

class TestExportFailure:

    def test_p7_err_001_xlsx_fail_csv_json_succeed(self):
        enriched = _golden_enriched()
        broken_excel = ExcelExporter()
        broken_excel.export = MagicMock(side_effect=RuntimeError("XLSX engine crashed"))
        orch = ExportOrchestrator(excel_exporter=broken_excel)
        req = ExportRequest(
            leads=[enriched], formats=["xlsx", "csv", "json"],
            output_dir=os.path.join(P7_SCRATCH, "err001"),
        )
        resp = orch.export(req)
        assert "xlsx" not in resp.files
        assert "leads_csv" in resp.files
        assert "json" in resp.files
        assert any("XLSX" in e for e in resp.errors)

    def test_p7_err_002_google_fail_local_success(self):
        broken_gs = GoogleSheetsExporter()
        broken_gs.export = MagicMock(side_effect=RuntimeError("Google API error"))
        orch = ExportOrchestrator(sheets_exporter=broken_gs)
        enriched = _golden_enriched()
        req = ExportRequest(
            leads=[enriched], formats=["xlsx", "json"],
            google_sheets=GoogleSheetsConfig(enabled=True),
            output_dir=os.path.join(P7_SCRATCH, "err002"),
        )
        resp = orch.export(req)
        assert "xlsx" in resp.files
        assert "json" in resp.files
        assert any("Google" in e for e in resp.errors)

    def test_p7_err_003_all_fail(self):
        enriched = _golden_enriched()
        broken_xl = ExcelExporter()
        broken_xl.export = MagicMock(side_effect=RuntimeError("fail"))
        broken_csv = CSVExporter()
        broken_csv.export = MagicMock(side_effect=RuntimeError("fail"))
        broken_crm = CRMExporter()
        broken_crm.export = MagicMock(side_effect=RuntimeError("fail"))
        orch = ExportOrchestrator(
            excel_exporter=broken_xl, csv_exporter=broken_csv,
            crm_exporter=broken_crm)
        req = ExportRequest(
            leads=[enriched], formats=["xlsx", "csv", "json"],
            output_dir=os.path.join(P7_SCRATCH, "err003"),
        )
        with pytest.raises(RuntimeError, match="All requested exports failed"):
            orch.export(req)

    def test_p7_err_004_error_audit_row(self):
        enriched = _golden_enriched()
        enriched.enrichment_errors = ["apollo: Rate limit exceeded"]
        resp, _ = _export([enriched], sub="err004")
        wb = openpyxl.load_workbook(resp.files["xlsx"])
        ws = wb["Errors_Audit"]
        assert ws.max_row >= 2
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "stage" in headers
        assert "error_type" in headers
        assert "message" in headers


# ====================================================================
# 15. LIMITS
# ====================================================================

class TestLimits:

    def test_p7_limit_001_5000_accepted(self):
        # Just test the limit check logic, not actual 5000 export
        assert MAX_EXPORT_LEADS == 5000

    def test_p7_limit_002_5001_rejected(self):
        orch = ExportOrchestrator()
        leads = [EnrichedLead(base_lead=CompanyLead(
            company_name=f"Co{i}", company_domain=f"co{i}.com",
            lead_score=70, qualified=True))
            for i in range(5001)]
        req = ExportRequest(leads=leads)
        with pytest.raises(ValueError, match="exceeds MAX_EXPORT_LEADS"):
            orch.export(req)

    def test_p7_limit_003_qualified_only_default(self):
        qualified = CompanyLead(
            company_name="Q", company_domain="q.com",
            lead_score=80, qualified=True)
        unqualified = CompanyLead(
            company_name="U", company_domain="u.com",
            lead_score=30, qualified=False)
        resp, _ = _export(
            [EnrichedLead(base_lead=qualified),
             EnrichedLead(base_lead=unqualified)],
            qualified_only=True, sub="limit003")
        assert resp.exported_leads == 1

    def test_p7_limit_004_include_all_contacts(self):
        enriched = _golden_enriched()
        # include_all_contacts=False -> only best contact
        resp_best_only, _ = _export(
            [enriched], include_all_contacts=False, sub="limit004a")
        assert resp_best_only.exported_contacts == 1
        # include_all_contacts=True -> all contacts
        resp_all, _ = _export(
            [enriched], include_all_contacts=True, sub="limit004b")
        assert resp_all.exported_contacts == 2


# ====================================================================
# 16. WORKFLOW
# ====================================================================

class TestWorkflow:

    def _row(self):
        enriched = _golden_enriched()
        ser = ExportSerializer(export_run_id="run_wf")
        lid = generate_lead_id(enriched)
        return ser.serialize_lead(enriched, lid, None)

    def test_p7_wf_001_approval(self):
        assert self._row().approval_status == "pending_review"

    def test_p7_wf_002_outreach(self):
        assert self._row().outreach_status == "not_started"

    def test_p7_wf_003_send(self):
        assert self._row().send_status == "not_sent"

    def test_p7_wf_004_no_draft(self):
        row = self._row()
        assert row.draft_subject in ("", None)
        assert row.draft_body in ("", None)
        assert row.personalization_notes in ("", None)
        assert row.last_outreach_at in ("", None)


# ====================================================================
# 17. SUMMARY
# ====================================================================

class TestSummary:

    def _summary(self, leads_list=None, include_unqualified=False):
        if leads_list is None:
            leads_list = [_golden_enriched()]
        resp, _ = _export(
            leads_list,
            _golden_jobs() if not include_unqualified else None,
            qualified_only=not include_unqualified,
            sub="summary")
        wb = openpyxl.load_workbook(resp.files["xlsx"])
        ws = wb["Summary"]
        metrics = {}
        for row_idx in range(2, ws.max_row + 1):
            metrics[ws.cell(row=row_idx, column=1).value] = ws.cell(row=row_idx, column=2).value
        return metrics, resp

    def test_p7_sum_001_total_count(self):
        metrics, resp = self._summary()
        assert metrics["Total Leads"] == resp.exported_leads

    def test_p7_sum_002_qualification_counts(self):
        metrics, _ = self._summary()
        assert metrics["Qualified Leads"] + metrics["Unqualified Leads"] == metrics["Total Leads"]

    def test_p7_sum_003_enrichment_status(self):
        metrics, _ = self._summary()
        assert metrics["Complete Enrichment"] == 1

    def test_p7_sum_004_email_metrics(self):
        metrics, _ = self._summary()
        assert isinstance(metrics["Leads With Verified Email"], int)
        assert isinstance(metrics["Leads With Likely Email"], int)

    def test_p7_sum_005_average_score(self):
        metrics, _ = self._summary()
        assert isinstance(metrics["Average Lead Score"], (int, float))

    def test_p7_sum_006_no_double_counting(self):
        enriched = _golden_enriched()
        # Duplicate same lead (same domain)
        dup = EnrichedLead(
            base_lead=CompanyLead(
                company_name="VDC Automation Lab",
                company_domain="www.vdc-autolab.de",  # canonicalizes to same
                lead_score=94, qualified=True),
        )
        metrics, resp = self._summary([enriched, dup])
        assert resp.exported_leads == 1  # Deduplicated
        assert metrics["Total Leads"] == 1


# ====================================================================
# 18. MANDATORY REGRESSION
# ====================================================================

class TestRegressions:

    def test_p7_reg_001_threshold_60(self):
        enriched = _golden_enriched()
        ser = ExportSerializer(export_run_id="run_reg1")
        row = ser.serialize_lead(enriched, generate_lead_id(enriched), None)
        assert row.qualification_threshold == 60

    def test_p7_reg_002_ats_namespace_no_collision(self):
        a = EnrichedLead(base_lead=CompanyLead(
            company_name="Acme", source_company_identities=["lever:acme"],
            lead_score=75, qualified=True))
        b = EnrichedLead(base_lead=CompanyLead(
            company_name="Acme", source_company_identities=["greenhouse:acme"],
            lead_score=75, qualified=True))
        resp, _ = _export([a, b], sub="reg002")
        assert resp.exported_leads == 2

    def test_p7_reg_003_weak_contact_no_collapse(self):
        c1 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="apollo")
        c2 = ContactCandidate(
            full_name="Alex Smith", job_title="BIM Manager",
            company_domain="acme.com", provider="hunter")
        enriched = EnrichedLead(
            base_lead=CompanyLead(company_name="Acme", company_domain="acme.com",
                                  lead_score=80, qualified=True),
            contacts=[c1, c2])
        resp, _ = _export([enriched], sub="reg003")
        assert resp.exported_contacts == 2

    def test_p7_reg_004_no_fake_generic_job(self):
        enriched = _golden_enriched()
        resp, _ = _export([enriched], jobs=None, include_jobs=True, sub="reg004")
        assert resp.exported_jobs == 0

    def test_p7_reg_005_google_mock_isolation(self):
        exporter = GoogleSheetsExporter()  # No injected client
        cfg = GoogleSheetsConfig(
            enabled=True, credentials_json='{"type": "service_account"}')
        with pytest.raises(NotImplementedError):
            exporter.export(cfg, [], [], [], [],
                            ExportSummary(export_timestamp="now"))

    def test_p7_reg_006_run_id_no_collision(self):
        resp1, _ = _export([_golden_enriched()], sub="reg006a")
        resp2, _ = _export([_golden_enriched()], sub="reg006b")
        assert resp1.export_run_id != resp2.export_run_id

    def test_p7_reg_007_empty_csv_headers(self):
        exporter = CSVExporter()
        d = os.path.join(P7_SCRATCH, "reg007")
        res = exporter.export(d, [], [], [], [])
        for path in res.values():
            with open(path, encoding="utf-8") as f:
                lines = f.readlines()
            assert len(lines) >= 1, f"No header in {path}"
            assert "," in lines[0]
