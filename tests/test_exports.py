import os
import shutil
import time
import pytest
import openpyxl
from datetime import datetime, timezone

from backend.schemas import CompanyLead, EnrichedLead, ContactCandidate, StructuredJob
from sales_engine.exports.schemas import (
    ExportRequest,
    GoogleSheetsConfig,
    LeadExportRow,
    ContactExportRow,
    JobExportRow,
    ErrorAuditRow,
    ExportSummary,
)
from sales_engine.exports.export_serializer import (
    ExportSerializer,
    generate_lead_id,
    generate_contact_id,
    clean_scalar_list,
)
from sales_engine.exports.excel_exporter import ExcelExporter
from sales_engine.exports.csv_exporter import CSVExporter
from sales_engine.exports.crm_exporter import CRMExporter
from sales_engine.exports.google_sheets_exporter import GoogleSheetsExporter, MockGoogleSheetClient
from sales_engine.exports.export_orchestrator import ExportOrchestrator, MAX_EXPORT_LEADS

TEST_OUTPUT_DIR = os.path.join(os.getcwd(), "tests", "scratch_export_test")

@pytest.fixture(autouse=True)
def cleanup_scratch():
    if os.path.exists(TEST_OUTPUT_DIR):
        shutil.rmtree(TEST_OUTPUT_DIR, ignore_errors=True)
    yield
    if os.path.exists(TEST_OUTPUT_DIR):
        shutil.rmtree(TEST_OUTPUT_DIR, ignore_errors=True)

def create_sample_lead(name="Acme BIM Corp", domain="acme-bim.com", score=85, qualified=True):
    base_lead = CompanyLead(
        company_name=name,
        company_name_normalized=name.lower(),
        company_domain=domain,
        source_company_keys=["acme-bim"],
        source_company_identities=["lever:acme-bim"],
        job_count=3,
        total_job_count=3,
        unique_job_count=3,
        relevant_job_count=2,
        job_titles=["BIM Director", "Revit API Developer"],
        locations=["Berlin, Germany"],
        technologies=["Revit", "Dynamo", "Python"],
        signals=[{"signal": "revit_api", "evidence_count": 2}],
        evidence=["Develop Revit C# plugins", "Automate Dynamo scripts"],
        fit_score=30,
        intent_score=30,
        recency_score=15,
        evidence_score=10,
        lead_score=score,
        qualification_threshold=60,
        qualified=qualified,
        scoring_reasons=["High fit BIM role", "Active Revit automation hiring"]
    )
    contact1 = ContactCandidate(
        first_name="Jane",
        last_name="Doe",
        full_name="Jane Doe",
        job_title="Head of BIM",
        seniority="Head",
        department="Digital Delivery",
        company_name=name,
        company_domain=domain,
        work_email="jane.doe@acme-bim.com",
        email_status="verified",
        email_confidence=95,
        linkedin_url="https://www.linkedin.com/in/janedoe-bim/",
        buyer_role_match="exact",
        contact_score=90,
        data_sources=["apollo"],
        is_best_contact=True,
    )
    contact2 = ContactCandidate(
        first_name="Bob",
        last_name="Smith",
        full_name="Bob Smith",
        job_title="BIM Coordinator",
        seniority="Mid",
        department="Engineering",
        company_name=name,
        company_domain=domain,
        work_email="bob.smith@acme-bim.com",
        email_status="likely",
        email_confidence=75,
        linkedin_url="https://www.linkedin.com/in/bobsmith/",
        buyer_role_match="relevant",
        contact_score=70,
        data_sources=["hunter"],
        is_best_contact=False,
    )
    enriched = EnrichedLead(
        base_lead=base_lead,
        buyer_roles_searched=["Head of BIM", "BIM Director"],
        contacts=[contact1, contact2],
        best_contact=contact1,
        enrichment_status="complete",
        providers_used=["apollo", "hunter"],
        enrichment_errors=[]
    )
    return enriched

# =========================================================================
# REQUIRED CONTRACT & REGRESSION TESTS (A through J)
# =========================================================================

def test_a_phase5_threshold_60_survives_export():
    """A. Phase 5 threshold 60 survives export."""
    lead = create_sample_lead(score=85, qualified=True)
    assert lead.base_lead.qualification_threshold == 60
    serializer = ExportSerializer(export_run_id="run_thresh")
    lid = generate_lead_id(lead)
    row = serializer.serialize_lead(lead, lid, None)
    assert row.qualification_threshold == 60
    assert row.qualified is True

def test_b_lever_and_greenhouse_acme_without_domains_remain_distinct_leads():
    """B. Lever acme and Greenhouse acme without domains remain distinct leads."""
    lead_lever = CompanyLead(
        company_name="Acme",
        source_company_keys=["acme"],
        source_company_identities=["lever:acme"],
        lead_score=75,
        qualified=True
    )
    lead_gh = CompanyLead(
        company_name="Acme",
        source_company_keys=["acme"],
        source_company_identities=["greenhouse:acme"],
        lead_score=75,
        qualified=True
    )

    id_lever = generate_lead_id(lead_lever)
    id_gh = generate_lead_id(lead_gh)

    assert id_lever != id_gh
    assert "lever:acme" in id_lever
    assert "greenhouse:acme" in id_gh

    # Export both
    req = ExportRequest(
        leads=[EnrichedLead(base_lead=lead_lever), EnrichedLead(base_lead=lead_gh)],
        formats=["json"],
        output_dir=os.path.join(TEST_OUTPUT_DIR, "test_b")
    )
    orchestrator = ExportOrchestrator()
    resp = orchestrator.export(req)
    assert resp.exported_leads == 2

def test_c_two_weak_identity_contacts_surviving_phase6_remain_two_export_rows():
    """C. Two weak-identity contacts surviving Phase 6 remain two export rows."""
    c1 = ContactCandidate(
        full_name="Alex Smith",
        job_title="BIM Manager",
        company_name="Acme",
        company_domain="acme.com",
        provider="apollo"
    )
    c2 = ContactCandidate(
        full_name="Alex Smith",
        job_title="BIM Manager",
        company_name="Acme",
        company_domain="acme.com",
        provider="hunter"
    )

    lead = create_sample_lead()
    lead.contacts = [c1, c2]

    req = ExportRequest(
        leads=[lead],
        formats=["json"],
        output_dir=os.path.join(TEST_OUTPUT_DIR, "test_c")
    )
    orchestrator = ExportOrchestrator()
    resp = orchestrator.export(req)
    assert resp.exported_contacts == 2

def test_d_same_email_produces_one_contact_identity():
    """D. Same email produces one contact identity."""
    lead_id = "domain:acme.com"
    c1 = ContactCandidate(work_email="jane.doe@acme.com", full_name="Jane Doe")
    c2 = ContactCandidate(work_email="JANE.DOE@ACME.COM", full_name="Jane Doe (Apollo)")
    assert generate_contact_id(c1, lead_id) == generate_contact_id(c2, lead_id)
    assert generate_contact_id(c1, lead_id) == "email:jane.doe@acme.com"

def test_e_same_linkedin_produces_one_contact_identity():
    """E. Same LinkedIn produces one contact identity."""
    lead_id = "domain:acme.com"
    c1 = ContactCandidate(linkedin_url="https://www.linkedin.com/in/janedoe/?ref=1")
    c2 = ContactCandidate(linkedin_url="https://linkedin.com/in/janedoe")
    assert generate_contact_id(c1, lead_id) == generate_contact_id(c2, lead_id)
    assert generate_contact_id(c1, lead_id) == "linkedin:linkedin.com/in/janedoe"

def test_f_jobs_are_not_fabricated_when_raw_jobs_are_absent():
    """F. Jobs are not fabricated when raw jobs are absent."""
    lead = create_sample_lead()
    req = ExportRequest(
        leads=[lead],
        jobs=None,  # No raw StructuredJob records supplied
        formats=["xlsx", "csv", "json"],
        include_jobs=True,
        output_dir=os.path.join(TEST_OUTPUT_DIR, "test_f")
    )
    orchestrator = ExportOrchestrator()
    resp = orchestrator.export(req)

    # Exported jobs count must be 0
    assert resp.exported_jobs == 0

    # Excel Jobs sheet must have 0 data rows (header only)
    wb = openpyxl.load_workbook(resp.files["xlsx"])
    ws_jobs = wb["Jobs"]
    assert ws_jobs.max_row == 1

    # Errors / Audit sheet must contain diagnostic
    ws_errors = wb["Errors_Audit"]
    assert ws_errors.max_row >= 2
    row2_vals = [ws_errors.cell(row=2, column=col).value for col in range(1, ws_errors.max_column + 1)]
    assert "jobs_not_provided" in row2_vals

def test_g_domainless_ats_leads_receive_raw_job_audit_records():
    """G. Domainless ATS leads can still receive their raw job audit records."""
    lead = CompanyLead(
        company_name="Innovate AEC",
        source_company_keys=["innovate-100"],
        source_company_identities=["lever:innovate-100"],
        lead_score=80,
        qualified=True
    )
    enriched = EnrichedLead(base_lead=lead)

    job = StructuredJob(
        company_name="Innovate AEC",
        source="lever",
        source_company_key="innovate-100",
        job_title="BIM Computational Designer",
        job_url="https://jobs.lever.co/innovate-100/123",
        technologies=["Revit", "Python"],
        relevant_signals=[]
    )

    req = ExportRequest(
        leads=[enriched],
        jobs=[job],
        formats=["json"],
        include_jobs=True,
        output_dir=os.path.join(TEST_OUTPUT_DIR, "test_g")
    )
    orchestrator = ExportOrchestrator()
    resp = orchestrator.export(req)

    assert resp.exported_jobs == 1

def test_h_google_credentials_cannot_accidentally_activate_mock_client():
    """H. Google credentials cannot accidentally activate MockGoogleSheetClient."""
    exporter = GoogleSheetsExporter()  # No injected test client
    cfg = GoogleSheetsConfig(enabled=True, credentials_json='{"type": "service_account"}')
    with pytest.raises(NotImplementedError) as exc_info:
        exporter.export(cfg, [], [], [], [], ExportSummary(export_timestamp="now"))
    assert "Live Google Sheets client adapter is not yet implemented" in str(exc_info.value)

def test_i_two_exports_in_same_second_receive_different_export_run_ids():
    """I. Two exports in the same second receive different export_run_id values."""
    req = ExportRequest(leads=[create_sample_lead()], formats=["json"], output_dir=os.path.join(TEST_OUTPUT_DIR, "test_i"))
    orchestrator = ExportOrchestrator()
    resp1 = orchestrator.export(req)
    resp2 = orchestrator.export(req)
    assert resp1.export_run_id != resp2.export_run_id

def test_j_empty_csv_exports_still_contain_headers():
    """J. Empty CSV exports still contain headers."""
    exporter = CSVExporter()
    csv_dir = os.path.join(TEST_OUTPUT_DIR, "test_j_csv")
    res = exporter.export(csv_dir, [], [], [], [])

    for k, path in res.items():
        assert os.path.exists(path)
        with open(path, encoding="utf-8") as f:
            content = f.read()
            assert len(content.strip()) > 0  # Headers written!
            assert "," in content

# =========================================================================
# PHASE 5 SCORING CONTRACT VALIDATION
# =========================================================================

def test_phase5_scoring_contract_acceptance():
    """
    Acceptance test using a REAL valid Phase 5-style CompanyLead.
    Asserts all Phase 5 contractual bounds are preserved through Phase 7 export.
    """
    base_lead = CompanyLead(
        company_name="Contract Validation Corp",
        company_domain="contract-valid.com",
        source_company_identities=["lever:contract-valid"],
        fit_score=30,       # <= 30
        intent_score=30,    # <= 30
        recency_score=20,   # <= 20
        evidence_score=14,  # <= 20
        lead_score=94,      # 30 + 30 + 20 + 14 = 94
        qualification_threshold=60,
        qualified=True,
    )
    enriched = EnrichedLead(base_lead=base_lead)

    # Validate Phase 5 contractual bounds
    assert base_lead.fit_score <= 30, f"fit_score {base_lead.fit_score} exceeds max 30"
    assert base_lead.intent_score <= 30, f"intent_score {base_lead.intent_score} exceeds max 30"
    assert base_lead.recency_score <= 20, f"recency_score {base_lead.recency_score} exceeds max 20"
    assert base_lead.evidence_score <= 20, f"evidence_score {base_lead.evidence_score} exceeds max 20"

    # Validate component sum
    component_sum = base_lead.fit_score + base_lead.intent_score + base_lead.recency_score + base_lead.evidence_score
    assert base_lead.lead_score == component_sum, (
        f"lead_score {base_lead.lead_score} != component sum {component_sum} "
        f"(fit={base_lead.fit_score}, intent={base_lead.intent_score}, "
        f"recency={base_lead.recency_score}, evidence={base_lead.evidence_score})"
    )

    # Validate threshold and qualified flag
    assert base_lead.qualification_threshold == 60
    assert base_lead.qualified == (base_lead.lead_score >= base_lead.qualification_threshold)

    # Now export and verify Phase 7 preserves all values unchanged
    serializer = ExportSerializer(export_run_id="run_contract")
    lid = generate_lead_id(enriched)
    row = serializer.serialize_lead(enriched, lid, None)

    assert row.fit_score == 30
    assert row.intent_score == 30
    assert row.recency_score == 20
    assert row.evidence_score == 14
    assert row.lead_score == 94
    assert row.qualification_threshold == 60
    assert row.qualified is True

    # Verify component sum is preserved
    exported_sum = row.fit_score + row.intent_score + row.recency_score + row.evidence_score
    assert row.lead_score == exported_sum

def test_phase5_scoring_contract_unqualified():
    """Edge case: a lead below threshold must export as unqualified with threshold=60."""
    base_lead = CompanyLead(
        company_name="Low Score Corp",
        company_domain="lowscore.com",
        fit_score=10,
        intent_score=10,
        recency_score=5,
        evidence_score=0,
        lead_score=25,
        qualification_threshold=60,
        qualified=False,
    )
    enriched = EnrichedLead(base_lead=base_lead)
    serializer = ExportSerializer(export_run_id="run_unq")
    lid = generate_lead_id(enriched)
    row = serializer.serialize_lead(enriched, lid, None)

    assert row.qualified is False
    assert row.qualification_threshold == 60
    assert row.lead_score == 25
    assert row.lead_score < row.qualification_threshold

# =========================================================================
# ADDITIONAL USABILITY TESTS
# =========================================================================

def test_privacy_exclusion():
    serializer = ExportSerializer(export_run_id="run_test")
    lead = create_sample_lead()
    l_row = serializer.serialize_lead(lead, "domain:acme-bim.com", "email:jane.doe@acme-bim.com")
    l_dict = l_row.model_dump()
    forbidden = ["personal_email", "personal_phone", "mobile_phone", "api_key", "password"]
    for k in forbidden:
        assert k not in l_dict

def test_xlsx_creation_and_openpyxl_usability():
    lead = create_sample_lead()
    serializer = ExportSerializer(export_run_id="run_100")
    lid = generate_lead_id(lead)
    bcid = generate_contact_id(lead.best_contact, lid)

    l_rows = [serializer.serialize_lead(lead, lid, bcid)]
    c_rows = serializer.serialize_contacts(lead, lid, bcid)
    j_rows = serializer.serialize_jobs(lead, lid)
    e_rows = serializer.serialize_errors(lead, lid)
    summary = serializer.calculate_summary(l_rows, c_rows, e_rows)

    os.makedirs(TEST_OUTPUT_DIR, exist_ok=True)
    xlsx_path = os.path.join(TEST_OUTPUT_DIR, "SalesAI_Export_test.xlsx")

    exporter = ExcelExporter()
    exported = exporter.export(xlsx_path, l_rows, c_rows, j_rows, e_rows, summary)
    assert os.path.exists(exported)

    wb = openpyxl.load_workbook(exported)
    expected_sheets = ["Leads", "Contacts", "Jobs", "Errors_Audit", "Summary"]
    assert wb.sheetnames == expected_sheets

    ws_leads = wb["Leads"]
    assert ws_leads.freeze_panes == "A2"
    assert ws_leads.auto_filter.ref is not None

def test_google_sheets_snapshot_and_upsert_mock():
    mock_client = MockGoogleSheetClient(spreadsheet_id="test-sheet-id")
    sheets_exporter = GoogleSheetsExporter(client=mock_client)

    lead = create_sample_lead(name="Delta Design", domain="deltadesign.de", score=92)
    serializer = ExportSerializer(export_run_id="run_sheet")
    lid = generate_lead_id(lead)
    bcid = generate_contact_id(lead.best_contact, lid)

    l_rows = [serializer.serialize_lead(lead, lid, bcid)]
    c_rows = serializer.serialize_contacts(lead, lid, bcid)
    summary = serializer.calculate_summary(l_rows, c_rows, [])

    # 1. Snapshot mode
    cfg_snap = GoogleSheetsConfig(enabled=True, mode="snapshot", spreadsheet_id="test-sheet-id")
    res_snap = sheets_exporter.export(cfg_snap, l_rows, c_rows, [], [], summary)
    assert res_snap["status"] == "success"
    assert len(mock_client.get_sheet_values("Leads")) == 2

    # 2. Upsert mode: Re-running with same dataset must NOT duplicate rows
    cfg_upsert = GoogleSheetsConfig(enabled=True, mode="upsert", spreadsheet_id="test-sheet-id")
    res_upsert1 = sheets_exporter.export(cfg_upsert, l_rows, c_rows, [], [], summary)
    assert res_upsert1["leads_updated"] == 1
    assert res_upsert1["leads_inserted"] == 0
    assert len(mock_client.get_sheet_values("Leads")) == 2

def test_max_export_leads_limit():
    orchestrator = ExportOrchestrator()
    oversized_leads = [create_sample_lead(name=f"Co {i}", domain=f"co{i}.com") for i in range(MAX_EXPORT_LEADS + 1)]
    req = ExportRequest(leads=oversized_leads)
    with pytest.raises(ValueError) as exc_info:
        orchestrator.export(req)
    assert "exceeds MAX_EXPORT_LEADS" in str(exc_info.value)
