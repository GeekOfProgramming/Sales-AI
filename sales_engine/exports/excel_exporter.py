import os
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from sales_engine.exports.schemas import (
    LeadExportRow,
    ContactExportRow,
    JobExportRow,
    ErrorAuditRow,
    ExportSummary,
)

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

DATA_FONT = Font(name="Calibri", size=10, color="000000")
DATA_ALIGNMENT = Alignment(vertical="top")
DATA_ALIGNMENT_WRAP = Alignment(vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)

class ExcelExporter:
    """Produces professionally formatted multi-sheet Excel workbooks using openpyxl."""

    def __init__(self):
        pass

    def export(
        self,
        output_path: str,
        lead_rows: List[LeadExportRow],
        contact_rows: List[ContactExportRow],
        job_rows: List[JobExportRow],
        error_rows: List[ErrorAuditRow],
        summary: ExportSummary,
    ) -> str:
        """Build Excel workbook and save to output_path."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        wb = openpyxl.Workbook()
        # Default sheet -> rename to Leads
        ws_leads = wb.active
        ws_leads.title = "Leads"
        self._populate_leads_sheet(ws_leads, lead_rows)

        ws_contacts = wb.create_sheet("Contacts")
        self._populate_contacts_sheet(ws_contacts, contact_rows)

        ws_jobs = wb.create_sheet("Jobs")
        self._populate_jobs_sheet(ws_jobs, job_rows)

        ws_errors = wb.create_sheet("Errors_Audit")
        self._populate_errors_sheet(ws_errors, error_rows)

        ws_summary = wb.create_sheet("Summary")
        self._populate_summary_sheet(ws_summary, summary)

        wb.save(output_path)
        return output_path

    def _style_headers(self, ws, num_cols: int):
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

    def _auto_fit_columns(self, ws, wrap_col_names: Optional[List[str]] = None):
        wrap_cols = set(wrap_col_names or [])
        header_names = {}
        for col_idx in range(1, ws.max_column + 1):
            cell_val = ws.cell(row=1, column=col_idx).value
            if cell_val:
                header_names[col_idx] = str(cell_val)

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            col_name = header_names.get(col_idx, "")
            is_wrap = col_name in wrap_cols

            max_len = len(col_name)
            # Sample up to 100 rows for performance
            for row_idx in range(2, min(ws.max_row + 1, 102)):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))

            if is_wrap:
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 20), 45)
            else:
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

        # Enable auto filters across all columns
        if ws.max_column > 0 and ws.max_row > 0:
            last_letter = get_column_letter(ws.max_column)
            ws.auto_filter.ref = f"A1:{last_letter}{ws.max_row}"

    def _populate_leads_sheet(self, ws, rows: List[LeadExportRow]):
        headers = [
            "export_run_id", "lead_id", "company_name", "company_domain",
            "industry", "country", "employee_count", "lead_score", "qualified",
            "fit_score", "intent_score", "recency_score", "evidence_score",
            "total_job_count", "relevant_job_count", "qualification_threshold",
            "lead_reasons", "lead_evidence", "enrichment_status", "providers_used",
            "enrichment_errors", "best_contact_id", "best_contact_name",
            "best_contact_title", "best_contact_email", "best_contact_email_status",
            "best_contact_email_confidence", "best_contact_linkedin",
            "best_contact_score", "buyer_role_match", "outreach_status",
            "approval_status", "send_status", "draft_subject", "draft_body",
            "personalization_notes", "last_outreach_at", "owner", "notes",
            "created_at", "updated_at"
        ]
        ws.append(headers)
        self._style_headers(ws, len(headers))

        wrap_cols = {"lead_reasons", "lead_evidence", "enrichment_errors", "notes"}

        for r in rows:
            row_data = [
                r.export_run_id, r.lead_id, r.company_name or "", r.company_domain or "",
                r.industry or "", r.country or "", r.employee_count if r.employee_count is not None else "",
                r.lead_score, r.qualified, r.fit_score, r.intent_score, r.recency_score, r.evidence_score,
                r.total_job_count, r.relevant_job_count, r.qualification_threshold,
                r.lead_reasons, r.lead_evidence, r.enrichment_status, r.providers_used,
                r.enrichment_errors, r.best_contact_id or "", r.best_contact_name or "",
                r.best_contact_title or "", r.best_contact_email or "", r.best_contact_email_status or "",
                r.best_contact_email_confidence if r.best_contact_email_confidence is not None else "",
                r.best_contact_linkedin or "", r.best_contact_score if r.best_contact_score is not None else "",
                r.buyer_role_match or "", r.outreach_status, r.approval_status, r.send_status,
                r.draft_subject or "", r.draft_body or "", r.personalization_notes or "",
                r.last_outreach_at or "", r.owner or "", r.notes or "",
                r.created_at, r.updated_at
            ]
            ws.append(row_data)
            row_idx = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if headers[col_idx - 1] in wrap_cols:
                    cell.alignment = DATA_ALIGNMENT_WRAP
                else:
                    cell.alignment = DATA_ALIGNMENT

        self._auto_fit_columns(ws, list(wrap_cols))

    def _populate_contacts_sheet(self, ws, rows: List[ContactExportRow]):
        headers = [
            "lead_id", "contact_id", "company_name", "company_domain",
            "first_name", "last_name", "full_name", "job_title", "seniority",
            "department", "work_email", "email_status", "email_confidence",
            "linkedin_url", "buyer_role_match", "contact_score", "data_sources",
            "provider_person_id", "is_best_contact", "contact_rank"
        ]
        ws.append(headers)
        self._style_headers(ws, len(headers))

        for r in rows:
            row_data = [
                r.lead_id, r.contact_id, r.company_name or "", r.company_domain or "",
                r.first_name or "", r.last_name or "", r.full_name or "",
                r.job_title or "", r.seniority or "", r.department or "",
                r.work_email or "", r.email_status or "", r.email_confidence or 0,
                r.linkedin_url or "", r.buyer_role_match or "", r.contact_score or 0,
                r.data_sources or "", r.provider_person_id or "",
                r.is_best_contact, r.contact_rank
            ]
            ws.append(row_data)
            row_idx = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = DATA_ALIGNMENT

        self._auto_fit_columns(ws)

    def _populate_jobs_sheet(self, ws, rows: List[JobExportRow]):
        headers = [
            "lead_id", "company_name", "company_domain", "job_title",
            "job_url", "source", "source_company_key", "location",
            "employment_type", "posted_date", "seniority", "remote_status",
            "technologies", "relevant_signals", "signal_evidence", "is_relevant"
        ]
        ws.append(headers)
        self._style_headers(ws, len(headers))

        wrap_cols = {"technologies", "relevant_signals", "signal_evidence"}

        for r in rows:
            row_data = [
                r.lead_id, r.company_name or "", r.company_domain or "",
                r.job_title or "", r.job_url or "", r.source or "",
                r.source_company_key or "", r.location or "",
                r.employment_type or "", r.posted_date or "",
                r.seniority or "", r.remote_status or "",
                r.technologies, r.relevant_signals, r.signal_evidence,
                r.is_relevant
            ]
            ws.append(row_data)
            row_idx = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                if headers[col_idx - 1] in wrap_cols:
                    cell.alignment = DATA_ALIGNMENT_WRAP
                else:
                    cell.alignment = DATA_ALIGNMENT

        self._auto_fit_columns(ws, list(wrap_cols))

    def _populate_errors_sheet(self, ws, rows: List[ErrorAuditRow]):
        headers = [
            "export_run_id", "lead_id", "company_name", "stage",
            "error_type", "provider", "message", "source_reference", "created_at"
        ]
        ws.append(headers)
        self._style_headers(ws, len(headers))

        for r in rows:
            row_data = [
                r.export_run_id, r.lead_id or "", r.company_name or "",
                r.stage, r.error_type, r.provider or "",
                r.message, r.source_reference or "", r.created_at
            ]
            ws.append(row_data)
            row_idx = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = DATA_ALIGNMENT_WRAP

        self._auto_fit_columns(ws, ["message"])

    def _populate_summary_sheet(self, ws, summary: ExportSummary):
        headers = ["Metric", "Value"]
        ws.append(headers)
        self._style_headers(ws, 2)

        metrics = [
            ("Total Leads", summary.total_leads),
            ("Qualified Leads", summary.qualified_leads),
            ("Unqualified Leads", summary.unqualified_leads),
            ("Complete Enrichment", summary.complete_enrichment),
            ("Partial Enrichment", summary.partial_enrichment),
            ("Provider Errors", summary.provider_errors),
            ("Leads With Best Contact", summary.leads_with_best_contact),
            ("Leads With Verified Email", summary.leads_with_verified_email),
            ("Leads With Likely Email", summary.leads_with_likely_email),
            ("Leads Without Work Email", summary.leads_without_work_email),
            ("Average Lead Score", summary.average_lead_score),
            ("Export Timestamp", summary.export_timestamp),
        ]

        for m_name, m_val in metrics:
            ws.append([m_name, m_val])
            row_idx = ws.max_row
            ws.cell(row=row_idx, column=1).font = Font(name="Calibri", size=10, bold=True)
            ws.cell(row=row_idx, column=1).border = THIN_BORDER
            ws.cell(row=row_idx, column=2).font = DATA_FONT
            ws.cell(row=row_idx, column=2).border = THIN_BORDER

        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 25
